"""markdown-reveal-pptx-api — butun loyiha bitta faylda."""
from __future__ import annotations

import io
import logging
import os
import re
import tempfile
import unicodedata
import xml.etree.ElementTree as ET
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional

from defusedxml import ElementTree as DefusedET
from fastapi import FastAPI, HTTPException, Request
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from jinja2 import Environment, select_autoescape
from markdown_it import MarkdownIt
from pydantic import BaseModel, Field, field_validator

# ============================================================================
# LOGGING
# ============================================================================
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger("reveal-pptx")

# ============================================================================
# CONFIG
# ============================================================================
VERSION = "1.0.0"
THEMES = [
    "black", "white", "league", "beige", "sky", "night",
    "serif", "simple", "solarized", "blood", "moon", "dracula",
]
ASPECT_RATIOS = ["16:9", "4:3"]
MAX_MARKDOWN_SIZE = int(os.getenv("MAX_MARKDOWN_SIZE", "200000"))
MAX_SVG_COUNT = int(os.getenv("MAX_SVG_COUNT", "20"))
MAX_SVG_SIZE = int(os.getenv("MAX_SVG_SIZE", "500000"))

# ============================================================================
# EXCEPTIONS
# ============================================================================
class InvalidSVGError(ValueError): pass
class UnsafeSVGError(ValueError): pass
class SVGNotFoundError(ValueError): pass
class ExportError(RuntimeError): pass

# ============================================================================
# UTILS
# ============================================================================
def slugify(value: str) -> str:
    if not value:
        return "presentation"
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^\w\s-]", "", value).strip().lower()
    return re.sub(r"[-\s]+", "-", value) or "presentation"


@contextmanager
def temp_workspace() -> Iterator[Path]:
    with tempfile.TemporaryDirectory(prefix="reveal_") as tmp:
        yield Path(tmp)


# ============================================================================
# SECURITY
# ============================================================================
EVENT_HANDLER_RE = re.compile(r"^on[a-z]+$", re.IGNORECASE)
DANGEROUS_TAGS = {"script", "iframe", "object", "embed", "applet", "foreignObject", "set"}
DANGEROUS_URI_SCHEMES = {"javascript:", "data:text/html", "vbscript:"}
DANGEROUS_ATTRS = {"src", "href", "xlink:href", "action", "formaction", "poster", "background"}


def _is_dangerous_uri(uri: str) -> bool:
    if not uri:
        return False
    normalized = re.sub(r"\s+", "", uri).lower()
    return any(normalized.startswith(s) for s in DANGEROUS_URI_SCHEMES)


def _local_name(tag: str) -> str:
    return tag.split("}", 1)[1] if "}" in tag else tag


# ============================================================================
# SCHEMAS (Pydantic)
# ============================================================================
class SVGAsset(BaseModel):
    name: str = Field(..., min_length=1, max_length=100)
    content: str = Field(..., min_length=1)

    @field_validator("name")
    @classmethod
    def check_name(cls, v: str) -> str:
        if not v.lower().endswith(".svg"):
            raise ValueError("SVG asset name must end with .svg")
        if any(c in v for c in ["..", "/", "\\"]):
            raise ValueError("Invalid SVG name")
        return v


class ConvertRequest(BaseModel):
    title: Optional[str] = Field(default="Untitled Presentation", max_length=200)
    author: Optional[str] = Field(default=None, max_length=200)
    theme: str = Field(default="black")
    aspect_ratio: str = Field(default="16:9")
    markdown: Optional[str] = Field(default=None, max_length=MAX_MARKDOWN_SIZE)
    svg_assets: Optional[List[SVGAsset]] = Field(default=None)

    @field_validator("theme")
    @classmethod
    def check_theme(cls, v: str) -> str:
        if v not in THEMES:
            raise ValueError("Unsupported theme")
        return v

    @field_validator("aspect_ratio")
    @classmethod
    def check_aspect(cls, v: str) -> str:
        if v not in ASPECT_RATIOS:
            raise ValueError("Unsupported aspect ratio")
        return v

    @field_validator("svg_assets")
    @classmethod
    def check_svg_count(cls, v):
        if v and len(v) > MAX_SVG_COUNT:
            raise ValueError(f"Too many SVG assets (max {MAX_SVG_COUNT})")
        return v


# ============================================================================
# SVG SERVICE
# ============================================================================
def validate_svg(content: str) -> None:
    if not content or not content.strip() or len(content) > MAX_SVG_SIZE:
        raise InvalidSVGError("Invalid SVG asset")
    lower = content.lower()
    if "<script" in lower or "<?php" in lower:
        raise UnsafeSVGError("Unsafe SVG content")
    try:
        root = DefusedET.fromstring(content.encode("utf-8"))
    except ET.ParseError:
        raise InvalidSVGError("Invalid SVG asset")
    if _local_name(root.tag) != "svg":
        raise InvalidSVGError("Invalid SVG asset")
    for el in root.iter():
        if _local_name(el.tag) in DANGEROUS_TAGS:
            raise UnsafeSVGError("Unsafe SVG content")
        for attr, val in list(el.attrib.items()):
            a = _local_name(attr).lower()
            if EVENT_HANDLER_RE.match(a):
                raise UnsafeSVGError("Unsafe SVG content")
            if a in {x.lower() for x in DANGEROUS_ATTRS} and _is_dangerous_uri(val):
                raise UnsafeSVGError("Unsafe SVG content")
            if a == "style" and ("expression(" in val.lower() or "javascript:" in val.lower()):
                raise UnsafeSVGError("Unsafe SVG content")


def clean_svg(content: str) -> str:
    """Xavfli atributlarni olib tashlaydi."""
    try:
        root = ET.fromstring(content.encode("utf-8"))
    except Exception:
        return content
    for el in root.iter():
        for attr in list(el.attrib.keys()):
            a = _local_name(attr).lower()
            if EVENT_HANDLER_RE.match(a):
                del el.attrib[attr]
            elif a in {x.lower() for x in DANGEROUS_ATTRS} and _is_dangerous_uri(el.attrib.get(attr, "")):
                del el.attrib[attr]
    buf = io.BytesIO()
    ET.ElementTree(root).write(buf, encoding="utf-8", xml_declaration=False)
    return buf.getvalue().decode("utf-8")


def process_svg_assets(assets: List[SVGAsset], workspace: Path) -> Dict[str, str]:
    svg_map: Dict[str, str] = {}
    seen = set()
    for a in assets:
        if a.name in seen:
            raise InvalidSVGError(f"Duplicate SVG asset: {a.name}")
        seen.add(a.name)
        validate_svg(a.content)
        cleaned = clean_svg(a.content)
        (workspace / a.name).write_text(cleaned, encoding="utf-8")
        svg_map[a.name] = cleaned
    return svg_map


def replace_svg_in_html(html: str, svg_map: Dict[str, str]) -> str:
    def _sub(m: re.Match) -> str:
        name = m.group(1)
        if name not in svg_map:
            raise SVGNotFoundError(f"SVG asset not found: {name}")
        svg = svg_map[name]
        if "<svg" in svg:
            svg = svg.replace("<svg", '<svg class="slide-svg" style="max-width:100%;height:auto;"', 1)
        return svg
    return re.sub(
        r'<img[^>]*src="__SVG__([^"]+)__"[^>]*/?>',
        _sub, html, flags=re.IGNORECASE
    )


# ============================================================================
# MARKDOWN SERVICE
# ============================================================================
def split_slides(markdown: str) -> List[str]:
    if not markdown:
        return []
    slides = re.split(r"\n\s*---\s*\n", markdown)
    return [s.strip() for s in slides if s.strip()]


def render_slide(md_text: str, svg_map: Dict[str, str]) -> str:
    # svg:file.svg → placeholder
    def _svg_sub(m: re.Match) -> str:
        alt, ref = m.group(1), m.group(2)
        if ref.startswith("svg:"):
            name = ref[4:]
            if name not in svg_map:
                raise SVGNotFoundError(f"SVG asset not found: {name}")
            return f'![{alt}](__SVG__{name}__)'
        return m.group(0)

    processed = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)", _svg_sub, md_text)
    md = MarkdownIt("commonmark", {"html": True, "linkify": True, "typographer": True})
    html = md.render(processed)

    # Fragmentlar: <li>...<!-- .element: class="fragment" -->
    html = re.sub(
        r"(<li[^>]*>)(.*?)(</li>)\s*<!--\s*\.element:\s*class=\"fragment\"\s*-->",
        lambda m: m.group(1)[:-1] + ' class="fragment">' + m.group(2) + m.group(3),
        html, flags=re.DOTALL
    )

    return replace_svg_in_html(html, svg_map)


def render_all_slides(slides: List[str], svg_map: Dict[str, str]) -> List[str]:
    return [render_slide(s, svg_map) for s in slides]


# ============================================================================
# REVEAL SERVICE (Jinja2 template inline)
# ============================================================================
REVEAL_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{{ title }}</title>
{% if author %}<meta name="author" content="{{ author }}">{% endif %}
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/dist/reset.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/dist/reveal.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/dist/theme/{{ theme }}.css">
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/plugin/highlight/monokai.css">
<style>
.reveal .slide-svg{max-width:90%;height:auto;display:block;margin:1em auto}
.reveal pre{font-size:.6em}.reveal code{font-family:'Fira Code',Consolas,monospace}
.reveal table{border-collapse:collapse;margin:1em auto}
.reveal th,.reveal td{border:1px solid #555;padding:6px 12px}
.reveal blockquote{border-left:4px solid #888;padding-left:1em;color:#ccc}
</style>
</head>
<body>
<div class="reveal"><div class="slides">
{% for slide in slides %}<section>{{ slide | safe }}</section>{% endfor %}
</div></div>
<script src="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/dist/reveal.js"></script>
<script src="https://cdn.jsdelivr.net/npm/reveal.js@4.6.1/plugin/highlight/highlight.js"></script>
<script>Reveal.initialize({hash:true,width:{{ w }},height:{{ h }},margin:0.08,plugins:[RevealHighlight]});</script>
</body></html>"""


def render_reveal_html(
    slides: List[str],
    title: str,
    author: Optional[str],
    theme: str,
    aspect_ratio: str,
    svg_map: Dict[str, str],
) -> str:
    rendered = render_all_slides(slides, svg_map)
    w, h = (1920, 1080) if aspect_ratio == "16:9" else (1024, 768)
    env = Environment(autoescape=select_autoescape(["html"]))
    template = env.from_string(REVEAL_TEMPLATE)
    return template.render(
        title=title or "Untitled Presentation",
        author=author or "",
        theme=theme,
        w=w, h=h,
        slides=rendered,
    )


# ============================================================================
# EXPORT SERVICE (python-pptx)
# ============================================================================
def export_pptx(
    slides: List[str],
    title: str,
    aspect_ratio: str,
    workspace: Path,
) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN

    sw = Inches(13.333) if aspect_ratio == "16:9" else Inches(10.0)
    sh = Inches(7.5)
    prs = Presentation()
    prs.slide_width = sw
    prs.slide_height = sh
    blank = prs.slide_layouts[6]

    md = MarkdownIt("commonmark", {"html": True})

    for slide_md in slides:
        slide = prs.slides.add_slide(blank)
        tokens = md.parse(slide_md)
        y = Inches(0.5)
        margin = Inches(0.6)
        max_w = sw - Emu(margin * 2)

        i = 0
        while i < len(tokens):
            tok = tokens[i]

            if tok.type == "heading":
                level = tok.tag[1] if tok.tag.startswith("h") else "1"
                size = {"1": 44, "2": 34, "3": 26}.get(level, 24)
                tb = slide.shapes.add_textbox(margin, y, max_w, Inches(0.8))
                p = tb.text_frame.paragraphs[0]
                p.text = tok.content
                p.font.size = Pt(size)
                p.font.bold = True
                p.alignment = PP_ALIGN.LEFT
                y += Inches(0.9)

            elif tok.type == "paragraph_open":
                # Keyingi inline tokenni olamiz
                if i + 1 < len(tokens) and tokens[i + 1].type == "inline":
                    text = tokens[i + 1].content
                    tb = slide.shapes.add_textbox(margin, y, max_w, Inches(0.5))
                    p = tb.text_frame.paragraphs[0]
                    p.text = text
                    p.font.size = Pt(20)
                    tb.text_frame.word_wrap = True
                    y += Inches(0.55)
                    i += 1  # inline tokenni o'tkazib yuboramiz

            elif tok.type == "bullet_list_open" or tok.type == "ordered_list_open":
                items = []
                for line in slide_md.split("\n"):
                    m = re.match(r"^\s*[-*+]\s+(.+)$", line) or re.match(r"^\s*\d+\.\s+(.+)$", line)
                    if m:
                        items.append(m.group(1).split("<!--")[0].strip())
                if items:
                    tb = slide.shapes.add_textbox(margin, y, max_w, Inches(len(items) * 0.45))
                    tf = tb.text_frame
                    tf.word_wrap = True
                    for idx, item in enumerate(items):
                        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                        p.text = f"• {item}"
                        p.font.size = Pt(20)
                    y += Inches(len(items) * 0.45 + 0.2)

            elif tok.type in ("fence", "code_block"):
                tb = slide.shapes.add_textbox(margin, y, max_w, Inches(2.0))
                p = tb.text_frame.paragraphs[0]
                p.text = tok.content
                p.font.size = Pt(14)
                p.font.name = "Consolas"
                tb.text_frame.word_wrap = True
                y += Inches(2.2)

            i += 1

    out = workspace / "presentation.pptx"
    prs.save(str(out))
    return out


# ============================================================================
# FASTAPI APP
# ============================================================================
app = FastAPI(title="markdown-reveal-pptx-api", version=VERSION)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(RequestValidationError)
async def _val_err(_: Request, exc: RequestValidationError):
    msg = exc.errors()[0]["msg"] if exc.errors() else "Validation error"
    return JSONResponse(status_code=422, content={"detail": msg})


@app.exception_handler(HTTPException)
async def _http_err(_: Request, exc: HTTPException):
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})


@app.exception_handler(Exception)
async def _any_err(_: Request, exc: Exception):
    logger.exception("Unhandled: %s", exc)
    return JSONResponse(status_code=500, content={"detail": "Internal server error"})


@app.get("/")
def health():
    return {"service": "markdown-reveal-pptx-api", "version": VERSION, "status": "ok"}


@app.get("/themes")
def themes():
    return {"themes": THEMES}


@app.post("/validate")
def validate(req: ConvertRequest):
    if not req.markdown or not req.markdown.strip():
        return {"valid": False, "errors": [{"field": "markdown", "message": "Markdown content is required"}]}
    errs = []
    if req.svg_assets:
        for a in req.svg_assets:
            try:
                validate_svg(a.content)
            except Exception as e:
                errs.append({"field": "svg_assets", "message": f"{a.name}: {e}"})
    if errs:
        return {"valid": False, "errors": errs}
    return {
        "valid": True,
        "slides_count": len(req.markdown.split("---")),
        "svg_assets_count": len(req.svg_assets or []),
    }


@app.post("/preview")
def preview(req: ConvertRequest):
    if not req.markdown or not req.markdown.strip():
        raise HTTPException(status_code=400, detail="Markdown content is required")
    try:
        with temp_workspace() as ws:
            svg_map = process_svg_assets(req.svg_assets or [], ws)
            slides = split_slides(req.markdown)
            html = render_reveal_html(
                slides=slides,
                title=req.title,
                author=req.author,
                theme=req.theme,
                aspect_ratio=req.aspect_ratio,
                svg_map=svg_map,
            )
            return {"html": html}
    except (InvalidSVGError, UnsafeSVGError, SVGNotFoundError) as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.exception("Preview failed: %s", e)
        raise HTTPException(status_code=500, detail="Failed to generate preview")


@app.post("/convert")
def convert(req: ConvertRequest):
    if not req.markdown or not req.markdown.strip():
        raise HTTPException(status_code=400, detail="Markdown content is required")
    try:
        with temp_workspace() as ws:
            svg_map = process_svg_assets(req.svg_assets or [], ws)
            slides = split_slides(req.markdown)
            pptx_path = export_pptx(
                slides=slides,
                title=req.title or "Presentation",
                aspect_ratio=req.aspect_ratio,
                workspace=ws,
            )
            filename = f"{slugify(req.title or 'presentation')}.pptx"
            return FileResponse(
                path=str(pptx_path),
                media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                filename=filename,
            )
    except (InvalidSVGError, UnsafeSVGError, SVGNotFoundError) as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.exception("Convert failed: %s", e)
        raise HTTPException(status_code=500, detail="Failed to export presentation")


# ============================================================================
# LOCAL RUN (agar to'g'ridan-to'g'ri ishga tushirilsa)
# ============================================================================
if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000, reload=False)